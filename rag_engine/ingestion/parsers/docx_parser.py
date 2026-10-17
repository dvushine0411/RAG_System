"""
rag_engine/ingestion/parsers/docx_parser.py
DOCX (Microsoft Word) parser using python-docx.
Extracts paragraphs and tables, retaining section headers when present.
"""

import io
from typing import Any, Dict, List
import docx


def parse_docx(file_bytes: bytes) -> List[Dict[str, Any]]:
    """
    Parses a DOCX byte stream and returns structured text segments.

    Args:
        file_bytes: Raw bytes of the DOCX document.

    Returns:
        List of dicts: [{"text": str, "metadata": {"section_title": str}}]
    """
    results: List[Dict[str, Any]] = []
    doc = docx.Document(io.BytesIO(file_bytes))

    current_section = "Document Header"
    current_buffer: List[str] = []

    def flush_buffer():
        nonlocal current_buffer
        combined_text = "\n".join(current_buffer).strip()
        if combined_text:
            results.append({
                "text": combined_text,
                "metadata": {"section_title": current_section}
            })
        current_buffer = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        # Check if style indicates a heading
        style_name = paragraph.style.name.lower() if paragraph.style and paragraph.style.name else ""
        if "heading" in style_name:
            flush_buffer()
            current_section = text
        else:
            current_buffer.append(text)

    # Process tables
    for table in doc.tables:
        table_rows: List[str] = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            if any(row_cells):
                table_rows.append(" | ".join(row_cells))
        if table_rows:
            table_text = "\n".join(table_rows)
            current_buffer.append(table_text)

    flush_buffer()
    return results
